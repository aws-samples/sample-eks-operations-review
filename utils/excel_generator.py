"""
Excel Export Generator - Detailed analysis results for filtering and analysis
"""
import pandas as pd
from datetime import datetime
from typing import Dict, Any, List
import io

class ExcelGenerator:
    """Generate comprehensive Excel reports for EKS analysis with multiple worksheets"""
    
    def __init__(self):
        self.workbook_data = {}
    
    def generate_report(self, analysis_results: Dict[str, Any], cluster_name: str) -> bytes:
        """Generate comprehensive Excel report and return as bytes"""
        
        # Collect all data for different worksheets
        all_checks = self._collect_all_checks(analysis_results)
        hardeneks_data = self._collect_hardeneks_data(analysis_results)
        dora_data = self._collect_dora_data(analysis_results)
        recommendations = self._collect_recommendations(analysis_results)
        compliance_data = self._collect_compliance_data(analysis_results)
        
        # Create Excel buffer
        buffer = io.BytesIO()
        
        # Create Excel writer with multiple sheets
        with pd.ExcelWriter(buffer, engine='openpyxl') as writer:
            
            # 1. Executive Summary Sheet
            self._create_executive_summary_sheet(writer, analysis_results, cluster_name)
            
            # 2. All Checks Sheet (Main data for analysis)
            if all_checks:
                checks_df = pd.DataFrame(all_checks)
                checks_df.to_excel(writer, sheet_name='All Checks', index=False)
                self._format_checks_sheet(writer.sheets['All Checks'], checks_df)
            
            # 3. HardenEKS Details Sheet
            if hardeneks_data:
                hardeneks_df = pd.DataFrame(hardeneks_data)
                hardeneks_df.to_excel(writer, sheet_name='HardenEKS Details', index=False)
                self._format_hardeneks_sheet(writer.sheets['HardenEKS Details'], hardeneks_df)
            
            # 4. DORA Compliance Sheet
            if dora_data:
                dora_df = pd.DataFrame(dora_data)
                dora_df.to_excel(writer, sheet_name='DORA Compliance', index=False)
                self._format_dora_sheet(writer.sheets['DORA Compliance'], dora_df)
            
            # 5. Recommendations Sheet
            if recommendations:
                rec_df = pd.DataFrame(recommendations)
                rec_df.to_excel(writer, sheet_name='Recommendations', index=False)
                self._format_recommendations_sheet(writer.sheets['Recommendations'], rec_df)
            
            # 6. Compliance Summary Sheet
            if compliance_data:
                comp_df = pd.DataFrame(compliance_data)
                comp_df.to_excel(writer, sheet_name='Compliance Summary', index=False)
                self._format_compliance_sheet(writer.sheets['Compliance Summary'], comp_df)
            
            # 7. Failed Checks Analysis Sheet
            failed_checks = [check for check in all_checks if check['Status'].upper() in ['FAIL', 'FAILED']]
            if failed_checks:
                failed_df = pd.DataFrame(failed_checks)
                failed_df.to_excel(writer, sheet_name='Failed Checks Analysis', index=False)
                self._format_failed_checks_sheet(writer.sheets['Failed Checks Analysis'], failed_df)
        
        buffer.seek(0)
        return buffer.getvalue()
    
    def _collect_all_checks(self, results: Dict[str, Any]) -> List[Dict]:
        """Collect all checks from different analyzers into a unified format"""
        all_checks = []
        
        # Basic Security Checks
        security_analysis = results.get('security_analysis', {})
        basic_checks = security_analysis.get('checks', [])
        
        for check in basic_checks:
            all_checks.append({
                'Check ID': check.get('id', 'N/A'),
                'Check Name': check.get('title', 'Unknown Check'),
                'Source': 'Basic Security',
                'Category': 'Security',
                'Status': check.get('status', 'Unknown'),
                'Severity': check.get('severity', 'N/A'),
                'Description': check.get('description', 'No description available'),
                'Finding': check.get('finding', check.get('description', 'No finding details')),
                'Risk Level': self._determine_risk_level(check.get('status'), check.get('severity')),
                'Recommendation': check.get('recommendation', 'See recommendations section'),
                'Command Used': check.get('command_used', 'aws eks describe-cluster --name [cluster]'),
                'AWS Documentation': self._get_aws_docs_link(check.get('title', '')),
                'Business Impact': self._get_business_impact(check.get('title', ''), check.get('status')),
                'Remediation Effort': self._estimate_effort(check.get('title', '')),
                'Compliance Frameworks': self._get_compliance_frameworks(check.get('title', ''))
            })
        
        # HardenEKS checks
        hardeneks_analysis = results.get('hardeneks_analysis', {})
        if not hardeneks_analysis:
            hardeneks_analysis = security_analysis.get('hardeneks_analysis', {})
        
        hardeneks_checks = hardeneks_analysis.get('checks', [])
        for check in hardeneks_checks:
            all_checks.append({
                'Check ID': check.get('check_id', check.get('id', 'N/A')),
                'Check Name': check.get('title', check.get('name', 'Unknown Check')),
                'Source': 'HardenEKS',
                'Category': check.get('category', 'HardenEKS'),
                'Status': check.get('status', 'Unknown'),
                'Severity': check.get('severity', check.get('priority', 'N/A')),
                'Description': check.get('description', 'No description available'),
                'Finding': check.get('finding', check.get('result', 'No finding details')),
                'Risk Level': self._determine_risk_level(check.get('status'), check.get('severity')),
                'Recommendation': check.get('recommendation', check.get('remediation', 'No recommendation')),
                'Command Used': check.get('command_used', check.get('check_command', 'N/A')),
                'AWS Documentation': check.get('documentation_url', 'https://aws.github.io/aws-eks-best-practices/'),
                'Business Impact': self._get_business_impact(check.get('title', ''), check.get('status')),
                'Remediation Effort': self._estimate_effort(check.get('title', '')),
                'Compliance Frameworks': 'CIS EKS Benchmark, AWS Best Practices'
            })
        
        # DORA Compliance checks
        dora_analysis = results.get('dora_analysis', {})
        if not dora_analysis:
            dora_analysis = security_analysis.get('dora_analysis', {})
        
        dora_checks = dora_analysis.get('checks', [])
        for check in dora_checks:
            all_checks.append({
                'Check ID': check.get('check_id', 'N/A'),
                'Check Name': check.get('title', 'Unknown Check'),
                'Source': 'DORA Compliance',
                'Category': 'DORA Compliance',
                'Status': check.get('status', 'Unknown'),
                'Severity': check.get('severity', 'N/A'),
                'Description': check.get('description', 'No description available'),
                'Finding': check.get('finding', 'No finding details'),
                'Risk Level': self._determine_risk_level(check.get('status'), check.get('severity')),
                'Recommendation': check.get('recommendation', check.get('guidance', 'No recommendation')),
                'Command Used': check.get('command_used', 'N/A'),
                'AWS Documentation': 'https://eur-lex.europa.eu/legal-content/EN/TXT/?uri=CELEX:32022R2554',
                'Business Impact': check.get('business_impact', 'Regulatory compliance impact'),
                'Remediation Effort': check.get('implementation_timeline', 'Medium'),
                'Compliance Frameworks': 'EU DORA Regulation 2022/2554'
            })
        
        return all_checks
    
    def _collect_hardeneks_data(self, results: Dict[str, Any]) -> List[Dict]:
        """Collect detailed HardenEKS analysis data"""
        security_analysis = results.get('security_analysis', {})
        hardeneks_analysis = results.get('hardeneks_analysis', {})
        if not hardeneks_analysis:
            hardeneks_analysis = security_analysis.get('hardeneks_analysis', {})
        
        if not hardeneks_analysis:
            return []
        
        hardeneks_data = []
        
        # Overall scores
        hardeneks_score = hardeneks_analysis.get('hardeneks_score', {})
        category_scores = hardeneks_score.get('category_scores', {})
        
        for category, cat_info in category_scores.items():
            hardeneks_data.append({
                'Category': category.replace('_', ' ').title(),
                'Score': f"{cat_info.get('score', 0):.1f}%",
                'Status': cat_info.get('status', 'Unknown'),
                'Checks Passed': cat_info.get('passed', 0),
                'Total Checks': cat_info.get('total_checks', 0),
                'Grade': self._get_grade_from_score(cat_info.get('score', 0)),
                'Key Issues': ', '.join([c.get('title', 'Unknown')[:50] for c in cat_info.get('checks', []) if c.get('status') == 'FAIL'][:3]),
                'Improvement Priority': self._get_priority_from_score(cat_info.get('score', 0))
            })
        
        return hardeneks_data
    
    def _collect_dora_data(self, results: Dict[str, Any]) -> List[Dict]:
        """Collect detailed DORA compliance data"""
        security_analysis = results.get('security_analysis', {})
        dora_analysis = results.get('dora_analysis', {})
        if not dora_analysis:
            dora_analysis = security_analysis.get('dora_analysis', {})
        
        if not dora_analysis:
            return []
        
        dora_data = []
        
        # Priority breakdown
        priority_breakdown = dora_analysis.get('priority_breakdown', {})
        
        for priority, data in priority_breakdown.items():
            dora_data.append({
                'Priority Level': priority,
                'Total Checks': data.get('total', 0),
                'Passed': data.get('passed', 0),
                'Failed': data.get('failed', 0),
                'Compliance Percentage': f"{data.get('compliance_percentage', 0):.1f}%",
                'Risk Level': data.get('risk_level', 'Unknown'),
                'Regulatory Impact': self._get_dora_regulatory_impact(data.get('compliance_percentage', 0)),
                'Action Required': 'Yes' if data.get('failed', 0) > 0 else 'No'
            })
        
        return dora_data
    
    def _collect_recommendations(self, results: Dict[str, Any]) -> List[Dict]:
        """Collect all recommendations from different sources"""
        recommendations = []
        
        # Basic security recommendations
        security_analysis = results.get('security_analysis', {})
        basic_recs = security_analysis.get('recommendations', [])
        
        for i, rec in enumerate(basic_recs, 1):
            recommendations.append({
                'ID': f"SEC-{i:03d}",
                'Title': rec.get('title', 'Unknown Recommendation'),
                'Priority': rec.get('priority', 'MEDIUM'),
                'Category': 'Security',
                'Description': rec.get('description', 'No description available'),
                'Implementation Command': rec.get('aws_cli', 'N/A'),
                'Expected Outcome': rec.get('expected_outcome', 'Improved security posture'),
                'Estimated Effort': self._estimate_effort(rec.get('title', '')),
                'Risk if Not Implemented': self._get_risk_if_not_implemented(rec.get('title', '')),
                'Compliance Impact': self._get_compliance_frameworks(rec.get('title', '')),
                'Dependencies': 'None',
                'Verification Steps': self._get_verification_steps(rec.get('title', ''))
            })
        
        # HardenEKS recommendations
        hardeneks_analysis = results.get('hardeneks_analysis', {})
        if not hardeneks_analysis:
            hardeneks_analysis = security_analysis.get('hardeneks_analysis', {})
        
        hardeneks_recs = hardeneks_analysis.get('recommendations', [])
        for i, rec in enumerate(hardeneks_recs, 1):
            recommendations.append({
                'ID': f"HEK-{i:03d}",
                'Title': rec.get('title', 'Unknown Recommendation'),
                'Priority': rec.get('priority', 'MEDIUM'),
                'Category': rec.get('category', 'HardenEKS'),
                'Description': rec.get('description', 'No description available'),
                'Implementation Command': rec.get('implementation_command', 'N/A'),
                'Expected Outcome': 'Enhanced security posture per EKS best practices',
                'Estimated Effort': self._estimate_effort(rec.get('title', '')),
                'Risk if Not Implemented': 'Security vulnerability may remain',
                'Compliance Impact': 'CIS EKS Benchmark, AWS Best Practices',
                'Dependencies': rec.get('dependencies', 'None'),
                'Verification Steps': rec.get('verification', 'Manual verification required')
            })
        
        # DORA recommendations
        dora_analysis = results.get('dora_analysis', {})
        if not dora_analysis:
            dora_analysis = security_analysis.get('dora_analysis', {})
        
        dora_recs = dora_analysis.get('recommendations', [])
        for i, rec in enumerate(dora_recs, 1):
            recommendations.append({
                'ID': f"DORA-{i:03d}",
                'Title': rec.get('title', 'Unknown Recommendation'),
                'Priority': rec.get('priority', 'MEDIUM'),
                'Category': rec.get('category', 'DORA Compliance'),
                'Description': rec.get('description', 'No description available'),
                'Implementation Command': 'N/A',
                'Expected Outcome': 'Improved DORA compliance',
                'Estimated Effort': rec.get('implementation_timeline', 'Medium'),
                'Risk if Not Implemented': 'Regulatory non-compliance',
                'Compliance Impact': 'EU DORA Regulation 2022/2554',
                'Dependencies': rec.get('dependencies', 'None'),
                'Verification Steps': 'Compliance audit verification'
            })
        
        return recommendations
    
    def _collect_compliance_data(self, results: Dict[str, Any]) -> List[Dict]:
        """Collect compliance framework summary data"""
        compliance_data = []
        
        # Add standard compliance frameworks with calculated scores
        frameworks = [
            ('CIS EKS Benchmark', 'v1.0.1', 'Security baseline for EKS'),
            ('NIST Cybersecurity Framework', 'v1.1', 'Comprehensive cybersecurity guidance'),
            ('SOC 2 Type II', '2017', 'Service organization controls'),
            ('EU DORA', '2024', 'Digital operational resilience'),
            ('PCI DSS', 'v3.2.1', 'Payment card industry security'),
            ('HIPAA Security Rule', '2013', 'Healthcare data protection'),
            ('ISO 27001', '2013', 'Information security management')
        ]
        
        security_analysis = results.get('security_analysis', {})
        total_checks = security_analysis.get('total_checks', 0)
        passed_checks = security_analysis.get('passed_checks', 0)
        
        # Calculate basic compliance percentage
        base_compliance = (passed_checks / total_checks * 100) if total_checks > 0 else 0
        
        for framework, version, description in frameworks:
            # Adjust compliance based on framework focus
            compliance_pct = self._adjust_compliance_for_framework(framework, base_compliance, results)
            
            compliance_data.append({
                'Framework': framework,
                'Version': version,
                'Description': description,
                'Compliance Percentage': f"{compliance_pct:.1f}%",
                'Status': self._get_compliance_status(compliance_pct),
                'Grade': self._get_grade_from_score(compliance_pct),
                'Gap Analysis': self._get_compliance_gaps(framework, compliance_pct),
                'Priority Actions': self._get_priority_actions(framework, compliance_pct),
                'Next Review Date': self._get_next_review_date()
            })
        
        return compliance_data
    
    def _create_executive_summary_sheet(self, writer, results: Dict[str, Any], cluster_name: str):
        """Create executive summary worksheet"""
        health = results.get('health_analysis', {})
        security = results.get('security_analysis', {})
        cluster_info = health.get('cluster_info', {})
        
        summary_data = {
            'Metric': [
                'Cluster Name', 'Analysis Date', 'Kubernetes Version', 'Total Security Checks',
                'Checks Passed', 'Checks Failed', 'Security Score', 'Overall Risk Level',
                'Critical Issues', 'Total Recommendations', 'Immediate Actions Required'
            ],
            'Value': [
                cluster_info.get('name', cluster_name),
                datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
                cluster_info.get('version', 'Unknown'),
                security.get('total_checks', 0),
                security.get('passed_checks', 0),
                security.get('failed_checks', 0),
                f"{(security.get('passed_checks', 0) / max(security.get('total_checks', 1), 1) * 100):.1f}%",
                self._get_overall_risk_level(security),
                len([c for c in security.get('checks', []) if c.get('severity') == 'HIGH' and c.get('status') == 'FAIL']),
                len(security.get('recommendations', [])),
                len([c for c in security.get('checks', []) if c.get('status') == 'FAIL' and c.get('severity') in ['HIGH', 'CRITICAL']])
            ]
        }
        
        summary_df = pd.DataFrame(summary_data)
        summary_df.to_excel(writer, sheet_name='Executive Summary', index=False)
        
        # Format the executive summary sheet
        worksheet = writer.sheets['Executive Summary']
        worksheet.column_dimensions['A'].width = 25
        worksheet.column_dimensions['B'].width = 30
    
    def _format_checks_sheet(self, worksheet, df):
        """Format the All Checks worksheet with filters and styling"""
        from openpyxl.styles import PatternFill, Font, Alignment
        from openpyxl.utils.dataframe import dataframe_to_rows
        
        # Set column widths
        column_widths = {
            'A': 12, 'B': 40, 'C': 15, 'D': 15, 'E': 12, 'F': 12, 
            'G': 50, 'H': 60, 'I': 12, 'J': 60, 'K': 40, 'L': 40,
            'M': 30, 'N': 15, 'O': 25
        }
        
        for col, width in column_widths.items():
            worksheet.column_dimensions[col].width = width
        
        # Add filters
        worksheet.auto_filter.ref = worksheet.dimensions
        
        # Style headers
        header_fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
        header_font = Font(color='FFFFFF', bold=True)
        
        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')
        
        # Color code status column (assuming Status is column E)
        status_col = 5  # Column E
        for row in range(2, len(df) + 2):
            cell = worksheet.cell(row=row, column=status_col)
            status = str(cell.value).upper()
            if status in ['FAIL', 'FAILED']:
                cell.fill = PatternFill(start_color='FFE6E6', end_color='FFE6E6', fill_type='solid')
                cell.font = Font(color='CC0000')
            elif status in ['PASS', 'PASSED']:
                cell.fill = PatternFill(start_color='E6F7E6', end_color='E6F7E6', fill_type='solid')
                cell.font = Font(color='006600')
            elif status in ['WARNING', 'WARN']:
                cell.fill = PatternFill(start_color='FFF2E6', end_color='FFF2E6', fill_type='solid')
                cell.font = Font(color='FF6600')
    
    def _format_hardeneks_sheet(self, worksheet, df):
        """Format HardenEKS Details sheet"""
        self._apply_basic_formatting(worksheet, df)
    
    def _format_dora_sheet(self, worksheet, df):
        """Format DORA Compliance sheet"""
        self._apply_basic_formatting(worksheet, df)
    
    def _format_recommendations_sheet(self, worksheet, df):
        """Format Recommendations sheet"""
        self._apply_basic_formatting(worksheet, df)
        
        # Set specific column widths for recommendations
        worksheet.column_dimensions['B'].width = 50  # Title
        worksheet.column_dimensions['E'].width = 70  # Description
        worksheet.column_dimensions['F'].width = 50  # Implementation Command
        
    def _format_compliance_sheet(self, worksheet, df):
        """Format Compliance Summary sheet"""
        self._apply_basic_formatting(worksheet, df)
    
    def _format_failed_checks_sheet(self, worksheet, df):
        """Format Failed Checks Analysis sheet with emphasis on critical items"""
        from openpyxl.styles import PatternFill, Font
        
        self._apply_basic_formatting(worksheet, df)
        
        # Highlight all rows as they're all failed checks
        failed_fill = PatternFill(start_color='FFE6E6', end_color='FFE6E6', fill_type='solid')
        for row in range(2, len(df) + 2):
            for col in range(1, len(df.columns) + 1):
                worksheet.cell(row=row, column=col).fill = failed_fill
    
    def _apply_basic_formatting(self, worksheet, df):
        """Apply basic formatting to worksheets"""
        from openpyxl.styles import PatternFill, Font, Alignment
        
        # Auto-adjust column widths
        for column in worksheet.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 70)
            worksheet.column_dimensions[column_letter].width = adjusted_width
        
        # Add filters
        worksheet.auto_filter.ref = worksheet.dimensions
        
        # Style headers
        header_fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
        header_font = Font(color='FFFFFF', bold=True)
        
        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')
    
    # Helper methods for data processing
    def _determine_risk_level(self, status: str, severity: str) -> str:
        """Determine overall risk level based on status and severity"""
        if not status:
            return 'Unknown'
            
        status = status.upper()
        severity = str(severity).upper() if severity else 'MEDIUM'
        
        if status in ['FAIL', 'FAILED']:
            if severity in ['HIGH', 'CRITICAL']:
                return 'Critical'
            elif severity == 'MEDIUM':
                return 'High'
            else:
                return 'Medium'
        elif status in ['WARNING', 'WARN']:
            return 'Low'
        else:
            return 'None'
    
    def _get_aws_docs_link(self, title: str) -> str:
        """Get relevant AWS documentation link based on check title"""
        title_lower = title.lower()
        
        if 'encryption' in title_lower:
            return 'https://docs.aws.amazon.com/eks/latest/userguide/encryption-at-rest.html'
        elif 'logging' in title_lower:
            return 'https://docs.aws.amazon.com/eks/latest/userguide/control-plane-logs.html'
        elif 'endpoint' in title_lower or 'api' in title_lower:
            return 'https://docs.aws.amazon.com/eks/latest/userguide/cluster-endpoint.html'
        elif 'network' in title_lower or 'vpc' in title_lower:
            return 'https://docs.aws.amazon.com/vpc/latest/userguide/VPC_SecurityGroups.html'
        elif 'iam' in title_lower or 'rbac' in title_lower:
            return 'https://docs.aws.amazon.com/eks/latest/userguide/security-iam.html'
        else:
            return 'https://docs.aws.amazon.com/eks/latest/userguide/security.html'
    
    def _get_business_impact(self, title: str, status: str) -> str:
        """Get business impact description"""
        if not status or status.upper() not in ['FAIL', 'FAILED']:
            return 'No immediate impact'
        
        title_lower = title.lower()
        
        if 'encryption' in title_lower:
            return 'Data breach risk, regulatory compliance violations'
        elif 'logging' in title_lower:
            return 'Limited incident response capability, compliance gaps'
        elif 'endpoint' in title_lower:
            return 'Increased attack surface, potential unauthorized access'
        elif 'network' in title_lower:
            return 'Lateral movement risk, data exfiltration potential'
        elif 'iam' in title_lower:
            return 'Privilege escalation risk, unauthorized access'
        else:
            return 'Security posture degradation, potential vulnerabilities'
    
    def _estimate_effort(self, title: str) -> str:
        """Estimate implementation effort"""
        title_lower = title.lower()
        
        if any(keyword in title_lower for keyword in ['encryption', 'logging', 'endpoint']):
            return 'Medium (2-4 hours)'
        elif any(keyword in title_lower for keyword in ['network', 'iam', 'rbac']):
            return 'High (1-2 days)'
        else:
            return 'Low (1-2 hours)'
    
    def _get_compliance_frameworks(self, title: str) -> str:
        """Get relevant compliance frameworks"""
        frameworks = []
        title_lower = title.lower()
        
        if 'encryption' in title_lower:
            frameworks.extend(['PCI DSS', 'HIPAA', 'SOC 2'])
        if 'logging' in title_lower:
            frameworks.extend(['SOC 2', 'NIST CSF'])
        if any(keyword in title_lower for keyword in ['network', 'endpoint', 'iam']):
            frameworks.extend(['CIS Benchmark', 'NIST CSF'])
        
        if not frameworks:
            frameworks = ['CIS Benchmark']
        
        return ', '.join(list(set(frameworks)))
    
    def _get_grade_from_score(self, score: float) -> str:
        """Convert numerical score to letter grade"""
        if score >= 90:
            return 'A'
        elif score >= 80:
            return 'B'
        elif score >= 70:
            return 'C'
        elif score >= 60:
            return 'D'
        else:
            return 'F'
    
    def _get_priority_from_score(self, score: float) -> str:
        """Get improvement priority based on score"""
        if score >= 80:
            return 'Low'
        elif score >= 60:
            return 'Medium'
        else:
            return 'High'
    
    def _get_dora_regulatory_impact(self, compliance_pct: float) -> str:
        """Get DORA regulatory impact assessment"""
        if compliance_pct >= 90:
            return 'Full Compliance'
        elif compliance_pct >= 70:
            return 'Minor Violations'
        else:
            return 'Significant Violations'
    
    def _adjust_compliance_for_framework(self, framework: str, base_compliance: float, results: Dict[str, Any]) -> float:
        """Adjust compliance percentage based on framework-specific factors"""
        # This is a simplified calculation - in reality, you'd map specific checks to frameworks
        adjustments = {
            'CIS EKS Benchmark': 0,
            'NIST Cybersecurity Framework': -5,
            'SOC 2 Type II': -3,
            'EU DORA': -10,  # More stringent
            'PCI DSS': -8,
            'HIPAA Security Rule': -7,
            'ISO 27001': -5
        }
        
        adjustment = adjustments.get(framework, 0)
        return max(0, min(100, base_compliance + adjustment))
    
    def _get_compliance_status(self, compliance_pct: float) -> str:
        """Get compliance status based on percentage"""
        if compliance_pct >= 80:
            return 'Compliant'
        elif compliance_pct >= 60:
            return 'Partially Compliant'
        else:
            return 'Non-Compliant'
    
    def _get_compliance_gaps(self, framework: str, compliance_pct: float) -> str:
        """Get compliance gap analysis"""
        if compliance_pct >= 80:
            return 'Minor gaps in advanced security controls'
        elif compliance_pct >= 60:
            return 'Moderate gaps in security implementation'
        else:
            return 'Major gaps requiring comprehensive remediation'
    
    def _get_priority_actions(self, framework: str, compliance_pct: float) -> str:
        """Get priority actions for compliance framework"""
        if compliance_pct >= 80:
            return 'Continue monitoring, address minor gaps'
        elif compliance_pct >= 60:
            return 'Implement medium-priority security controls'
        else:
            return 'Urgent: Address critical security vulnerabilities'
    
    def _get_next_review_date(self) -> str:
        """Get next compliance review date"""
        from datetime import datetime, timedelta
        next_review = datetime.now() + timedelta(days=90)
        return next_review.strftime('%Y-%m-%d')
    
    def _get_overall_risk_level(self, security: Dict[str, Any]) -> str:
        """Get overall risk level based on security analysis"""
        total_checks = security.get('total_checks', 0)
        failed_checks = security.get('failed_checks', 0)
        
        if total_checks == 0:
            return 'Unknown'
        
        failure_rate = failed_checks / total_checks
        
        if failure_rate >= 0.5:
            return 'Critical'
        elif failure_rate >= 0.3:
            return 'High'
        elif failure_rate >= 0.1:
            return 'Medium'
        else:
            return 'Low'
    
    def _get_risk_if_not_implemented(self, title: str) -> str:
        """Get risk description if recommendation not implemented"""
        title_lower = title.lower()
        
        if 'encryption' in title_lower:
            return 'High: Data exposure in case of breach'
        elif 'logging' in title_lower:
            return 'Medium: Limited visibility and audit capabilities'
        elif 'endpoint' in title_lower:
            return 'High: Increased attack surface exposure'
        elif 'network' in title_lower:
            return 'Medium: Potential lateral movement vulnerabilities'
        else:
            return 'Medium: General security posture degradation'
    
    def _get_verification_steps(self, title: str) -> str:
        """Get verification steps for recommendation"""
        title_lower = title.lower()
        
        if 'encryption' in title_lower:
            return 'Verify KMS key configuration and encryption status'
        elif 'logging' in title_lower:
            return 'Check CloudWatch logs and log group configuration'
        elif 'endpoint' in title_lower:
            return 'Verify API endpoint access configuration'
        elif 'network' in title_lower:
            return 'Test network policies and security group rules'
        else:
            return 'Manual verification of configuration changes'
